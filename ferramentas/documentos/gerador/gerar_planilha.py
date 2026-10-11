"""Gera a planilha padrão SafetySoluções (Inventário de Riscos + Plano de Ação 5W2H + Matriz 5x5).
Uso: python3 gerar_planilha.py dados.json saida.xlsx      (com dados)
     python3 gerar_planilha.py - saida.xlsx               (modelo em branco)
As colunas são as mesmas do inventário do PGR, então o sistema gera os dois a partir do mesmo dados.json."""
import json, sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.utils import get_column_letter as L

NAVY, LIGHT = '0F2A44', 'EEF2F6'
# Classificação (mesma da planilha original): nível = Probabilidade x Severidade
NIVEIS = [(1, 3, 'TRIVIAL', '00B0F0', 'Nenhuma ação necessária.'),
          (4, 8, 'TOLERÁVEL', '92D050', 'Nenhum controle adicional necessário.'),
          (9, 12, 'MODERADO', 'FFFF00', 'Controle adicional, se possível/viável.'),
          (13, 15, 'SUBSTANCIAL', 'FFC000', 'Controle necessário.'),
          (16, 25, 'INTOLERÁVEL', 'C00000', 'Ação imediata ou interrupção da atividade.')]
STATUS = [('ATRASADO', 'FF7C80'), ('EM ANDAMENTO', 'FFFF00'), ('PROGRAMADO', '9DC3E6'), ('REALIZADO', '92D050'), ('REPROGRAMADO', 'FFC000')]
MESES = ['JAN', 'FEV', 'MAR', 'ABR', 'MAI', 'JUN', 'JUL', 'AGO', 'SET', 'OUT', 'NOV', 'DEZ']

fino = Side(style='thin', color='9AA8B5')
BORDA = Border(left=fino, right=fino, top=fino, bottom=fino)
CENTRO = Alignment(horizontal='center', vertical='center', wrap_text=True)
ESQ = Alignment(horizontal='left', vertical='center', wrap_text=True)
fill = lambda c: PatternFill('solid', start_color=c, end_color=c)


def cab(ws, r, c, txt, r2=None, c2=None, cor=NAVY):
    if r2 or c2:
        ws.merge_cells(start_row=r, start_column=c, end_row=r2 or r, end_column=c2 or c)
    for rr in range(r, (r2 or r) + 1):
        for cc in range(c, (c2 or c) + 1):
            x = ws.cell(rr, cc)
            x.fill, x.border, x.alignment = fill(cor), BORDA, CENTRO
            x.font = Font(bold=True, color='FFFFFF', size=9)
    ws.cell(r, c).value = txt


def titulo(ws, txt, ult_col, emp):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ult_col)
    ws['A1'] = txt
    ws['A1'].font, ws['A1'].alignment = Font(bold=True, size=14, color=NAVY), CENTRO
    ws.row_dimensions[1].height = 26
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ult_col)
    ws['A2'] = f"Empresa: {emp.get('RAZAO_SOCIAL', '')}    CNPJ: {emp.get('CNPJ', '')}    Atualização: {emp.get('DATA_EMISSAO', '')}    Responsável: {emp.get('RESPONSAVEL_NOME', '')}"
    ws['A2'].font, ws['A2'].alignment = Font(size=9, color='5B6770'), CENTRO


def inventario(wb, dados, emp):
    ws = wb.active
    ws.title = 'INVENTÁRIO DE RISCOS'
    titulo(ws, 'INVENTÁRIO DE RISCOS OCUPACIONAIS – NR-01 (item 1.5.7.3)', 20, emp)
    heads = ['DESCRIÇÃO DO AMBIENTE', 'SETOR', 'GES', 'FUNÇÃO', 'PROCESSO / ATIVIDADES', 'AGENTE', 'PERIGO', 'FONTE / CIRCUNSTÂNCIA',
             'POSSÍVEIS LESÕES OU AGRAVOS À SAÚDE', 'LIMITE DE EXPOSIÇÃO', 'INTENS. / CONC.', 'TIPO / TEMPO DE EXPOSIÇÃO', 'TÉCNICA UTILIZADA']
    for i, h in enumerate(heads, 1):
        cab(ws, 4, i, h, r2=5)
    cab(ws, 4, 14, 'MEDIDAS DE PREVENÇÃO EXISTENTES', c2=16)
    for i, h in enumerate(['EPC / PROCEDIMENTO ADMINISTRATIVO', 'EPI', 'ATENUAÇÃO / FATOR DE PROTEÇÃO'], 14):
        cab(ws, 5, i, h)
    for i, h in enumerate(['PROBABILIDADE', 'SEVERIDADE', 'NÍVEL (P x S)', 'CLASSIFICAÇÃO DO RISCO'], 17):
        cab(ws, 4, i, h, r2=5)
    larg = [34, 14, 7, 26, 32, 13, 22, 22, 30, 11, 11, 13, 16, 34, 16, 13, 9, 9, 9, 15]
    for i, w in enumerate(larg, 1):
        ws.column_dimensions[L(i)].width = w
    ws.row_dimensions[4].height, ws.row_dimensions[5].height = 22, 36

    chaves = ['ambiente', 'setor', 'ghe', 'funcao', 'processo', 'agente', 'perigo', 'fonte', 'lesoes', 'limite',
              'intensidade', 'exposicao', 'tecnica', 'medidas', 'epi', 'atenuacao', 'probabilidade', 'severidade']
    linhas = dados or [dict.fromkeys(chaves, '') for _ in range(40)]
    r0 = 6
    for k, reg in enumerate(linhas):
        r = r0 + k
        for i, ch in enumerate(chaves, 1):
            v = reg.get(ch, '')
            c = ws.cell(r, i, v if v not in (None,) else '')
            c.border, c.font = BORDA, Font(size=9)
            c.alignment = CENTRO if i in (2, 3, 6, 10, 11, 12, 17, 18) else ESQ
        ws.cell(r, 19, f'=IF(OR(Q{r}="",R{r}=""),"",Q{r}*R{r})')
        f = f'=IF(S{r}="","",' + ''.join(f'IF(S{r}<={b},"{n}",' for _, b, n, _, _ in NIVEIS[:-1]) + f'"{NIVEIS[-1][2]}"' + ')' * len(NIVEIS)
        ws.cell(r, 20, f)
        for i in (19, 20):
            ws.cell(r, i).border, ws.cell(r, i).alignment, ws.cell(r, i).font = BORDA, CENTRO, Font(size=9, bold=True)
    ult = r0 + len(linhas) - 1

    # mescla as colunas de grupo (ambiente, setor, GES, função, processo) e o agente, como na planilha original
    if dados:
        def mescla(col, chave_fn):
            ini = r0
            for r in range(r0 + 1, ult + 2):
                if r > ult or chave_fn(linhas[r - r0]) != chave_fn(linhas[ini - r0]):
                    if r - 1 > ini:
                        ws.merge_cells(start_row=ini, start_column=col, end_row=r - 1, end_column=col)
                    ini = r
        g = lambda x: (x['ghe'], x['setor'], x['funcao'])
        for col, ch in ((1, 'ambiente'), (2, 'setor'), (3, 'ghe'), (4, 'funcao'), (5, 'processo')):
            mescla(col, lambda x, ch=ch: (g(x), x[ch]))
        mescla(6, lambda x: (g(x), x['agente']))

    faixa = f'T{r0}:T{ult}'
    for _, _, n, cor, _ in NIVEIS:
        ws.conditional_formatting.add(faixa, CellIsRule(operator='equal', formula=[f'"{n}"'], fill=fill(cor), font=Font(bold=True, color='FFFFFF' if n == 'INTOLERÁVEL' else '000000')))
    dv = DataValidation(type='list', formula1='"Físico,Químico,Biológico,Ergonômico,Acidente"', allow_blank=True)
    dv2 = DataValidation(type='whole', operator='between', formula1='1', formula2='5', allow_blank=True, error='Use valores de 1 a 5', showErrorMessage=True)
    dv3 = DataValidation(type='list', formula1='"Permanente,Intermitente,Eventual"', allow_blank=True)
    for v in (dv, dv2, dv3):
        ws.add_data_validation(v)
    dv.add(f'F{r0}:F{ult}'); dv2.add(f'Q{r0}:R{ult}'); dv3.add(f'L{r0}:L{ult}')
    ws.freeze_panes = 'D6'
    ws.page_setup.orientation, ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 'landscape', 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = '4:5'


def plano(wb, dados, emp):
    ws = wb.create_sheet('PLANO DE AÇÃO 5W2H')
    titulo(ws, 'PLANO DE AÇÃO 5W2H – NR-01 (item 1.5.5.2)', 23, emp)
    for i, h in enumerate(['SETOR', 'GES', 'O QUÊ? (MEDIDA DE PREVENÇÃO)', 'POR QUÊ?', 'COMO?', 'ONDE?', 'QUANTO?', 'PRIORIDADE'], 1):
        cab(ws, 4, i, h, r2=5)
    cab(ws, 4, 9, f"CRONOGRAMA {emp.get('ANO', '')}".strip(), c2=20)
    for i, m in enumerate(MESES, 9):
        cab(ws, 5, i, m)
    cab(ws, 4, 21, 'RESPONSÁVEL', r2=5)
    cab(ws, 4, 22, 'STATUS', r2=5)
    cab(ws, 4, 23, 'ACOMPANHAMENTO E AFERIÇÃO DE RESULTADOS', r2=5)
    for i, w in enumerate([14, 6, 32, 28, 30, 16, 16, 11] + [5] * 12 + [16, 15, 30], 1):
        ws.column_dimensions[L(i)].width = w
    ws.row_dimensions[5].height = 20
    linhas = dados or [{} for _ in range(20)]
    r0 = 6
    for k, a in enumerate(linhas):
        r = r0 + k
        vals = [a.get('setor', ''), a.get('ghe', ''), a.get('oque', ''), a.get('porque', ''), a.get('como', ''), a.get('onde', ''), a.get('quanto', ''), a.get('prioridade', '')]
        vals += ['X' if m in a.get('meses', []) else '' for m in MESES]
        vals += [a.get('responsavel', ''), a.get('status', ''), a.get('afericao', '')]
        for i, v in enumerate(vals, 1):
            c = ws.cell(r, i, v)
            c.border, c.font = BORDA, Font(size=9, bold=(9 <= i <= 20))
            c.alignment = ESQ if i in (3, 4, 5, 6, 7, 23) else CENTRO
        ws.row_dimensions[r].height = 48
    ult = r0 + len(linhas) - 1
    ws.conditional_formatting.add(f'I{r0}:T{ult}', CellIsRule(operator='equal', formula=['"X"'], fill=fill(NAVY), font=Font(bold=True, color='FFFFFF')))
    for n, cor in STATUS:
        ws.conditional_formatting.add(f'V{r0}:V{ult}', CellIsRule(operator='equal', formula=[f'"{n}"'], fill=fill(cor)))
    dv = DataValidation(type='list', formula1='"' + ','.join(n for n, _ in STATUS) + '"', allow_blank=True)
    dvp = DataValidation(type='list', formula1='"ALTA,MÉDIA,BAIXA"', allow_blank=True)
    ws.add_data_validation(dv); ws.add_data_validation(dvp)
    dv.add(f'V{r0}:V{ult}'); dvp.add(f'H{r0}:H{ult}')
    lg = ult + 2
    ws.cell(lg, 1, 'LEGENDA DO STATUS').font = Font(bold=True, size=9, color=NAVY)
    for i, (n, cor) in enumerate(STATUS):
        c = ws.cell(lg + 1 + i, 1, n)
        c.fill, c.font, c.border = fill(cor), Font(size=9, bold=True), BORDA
        ws.merge_cells(start_row=lg + 1 + i, start_column=1, end_row=lg + 1 + i, end_column=2)
    ws.freeze_panes = 'C6'
    ws.page_setup.orientation = 'landscape'
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0


def matriz(wb):
    ws = wb.create_sheet('MATRIZ DE RISCO 5X5')
    ws.merge_cells('A1:H1'); ws['A1'] = 'MATRIZ DE RISCO 5x5 – NÍVEL DE RISCO = PROBABILIDADE x SEVERIDADE'
    ws['A1'].font, ws['A1'].alignment = Font(bold=True, size=13, color=NAVY), CENTRO
    cab(ws, 3, 4, 'SEVERIDADE', c2=8)
    for i, s in enumerate(['Leve (1)', 'Baixa (2)', 'Moderada (3)', 'Alta (4)', 'Extrema (5)']):
        cab(ws, 4, 4 + i, s, cor='1B3F63')
    cab(ws, 5, 1, 'PROBABILIDADE', r2=9)
    for k, (n, p) in enumerate([('Muito Provável', 5), ('Provável', 4), ('Possível', 3), ('Pouco Provável', 2), ('Rara', 1)]):
        cab(ws, 5 + k, 2, n, cor='1B3F63'); cab(ws, 5 + k, 3, p, cor='1B3F63')
        for s in range(1, 6):
            v = p * s
            cor = next(c for a, b, _, c, _ in NIVEIS if a <= v <= b)
            c = ws.cell(5 + k, 3 + s, v)
            c.fill, c.alignment, c.border = fill(cor), CENTRO, BORDA
            c.font = Font(bold=True, size=12, color='FFFFFF' if cor == 'C00000' else '000000')
        ws.row_dimensions[5 + k].height = 28
    cab(ws, 12, 1, 'NÍVEL DO RISCO', c2=2); cab(ws, 12, 3, 'CLASSIFICAÇÃO', c2=4); cab(ws, 12, 5, 'TRATAMENTO NO PLANO DE AÇÃO', c2=8)
    for k, (a, b, n, cor, acao) in enumerate(NIVEIS):
        r = 13 + k
        for c1, c2, v, f in ((1, 2, f'{a} a {b}', cor), (3, 4, n, cor), (5, 8, acao, None)):
            ws.merge_cells(start_row=r, start_column=c1, end_row=r, end_column=c2)
            ws.cell(r, c1, v)
            for cc in range(c1, c2 + 1):
                x = ws.cell(r, cc)
                x.border, x.alignment = BORDA, CENTRO if f else ESQ
                if f: x.fill = fill(f)
            ws.cell(r, c1).font = Font(bold=bool(f), size=10, color='FFFFFF' if f == 'C00000' else '000000')
    for i, w in enumerate([14, 16, 6, 14, 14, 14, 14, 14], 1):
        ws.column_dimensions[L(i)].width = w


if __name__ == '__main__':
    src, dst = sys.argv[1], sys.argv[2]
    d = json.load(open(src, encoding='utf-8')) if src != '-' else {}
    emp = d.get('empresa', {})
    wb = Workbook()
    inventario(wb, d.get('inventario'), emp)
    plano(wb, d.get('plano'), emp)
    matriz(wb)
    wb.save(dst)
    print('ok', dst)
