"""Lê a planilha de Inventário de Riscos + Plano de Ação (.xlsm/.xlsx) e gera dados.json.
Uso: python3 extrair_planilha.py planilha.xlsm dados.json
O dados.json é a fonte única: dele saem a planilha padrão e o PGR."""
import json, re, sys
import openpyxl

INV_COLS = ['ambiente', 'setor', 'ghe', 'funcao', 'processo', 'agente', 'perigo', 'fonte', 'lesoes', 'limite',
            'intensidade', 'exposicao', 'tecnica', 'medidas', 'epi', 'atenuacao', 'probabilidade', 'severidade']
GRUPO = INV_COLS[:6]  # colunas que valem para o grupo (repetem nas linhas de baixo)


def limpa(v):
    if v is None:
        return ''
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    s = str(v).strip()
    s = re.sub(r'[ \t]+', ' ', s)
    s = re.sub(r'\s*\n\s*', '\n', s)
    return s


def valores(ws):
    """Matriz de valores com as células mescladas preenchidas."""
    mapa = {}
    for rg in ws.merged_cells.ranges:
        v = ws.cell(rg.min_row, rg.min_col).value
        for r in range(rg.min_row, rg.max_row + 1):
            for c in range(rg.min_col, rg.max_col + 1):
                mapa[(r, c)] = v
    return lambda r, c: mapa.get((r, c), ws.cell(r, c).value)


def acha_aba(wb, *palavras):
    for ws in wb.worksheets:
        n = ws.title.upper()
        if all(p in n for p in palavras):
            return ws
    return None


def ler_inventario(ws):
    val = valores(ws)
    # linha do cabeçalho = a que tem "SETOR" na coluna B
    cab = next(r for r in range(1, 30) if limpa(val(r, 2)).upper().startswith('SETOR'))
    linhas, ant = [], {}
    for r in range(cab + 2, ws.max_row + 1):
        reg = {k: limpa(val(r, i + 1)) for i, k in enumerate(INV_COLS)}
        if not any(reg[k] for k in INV_COLS[6:16]):
            continue
        for k in GRUPO:
            if not reg[k]:
                reg[k] = ant.get(k, '')
        for k in ('probabilidade', 'severidade'):
            try:
                reg[k] = int(float(reg[k]))
            except ValueError:
                reg[k] = None
        ag = reg['agente'].capitalize()
        reg['agente'] = {'Acidentes': 'Acidente', 'Fisico': 'Físico', 'Quimico': 'Químico', 'Biologico': 'Biológico', 'Ergonomico': 'Ergonômico'}.get(ag, ag)
        ant = reg
        linhas.append(reg)
    return linhas


MESES = ['JAN', 'FEV', 'MAR', 'ABR', 'MAI', 'JUN', 'JUL', 'AGO', 'SET', 'OUT', 'NOV', 'DEZ']


def ler_plano(ws):
    val = valores(ws)
    cab = next(r for r in range(1, 20) if limpa(val(r, 1)).upper().startswith('SETOR'))
    cols = {limpa(val(cab, c)).upper(): c for c in range(1, ws.max_column + 1) if val(cab, c)}
    col_mes = {m: cols[m] for m in MESES if m in cols}
    col_resp = next((c for c in range(1, ws.max_column + 1) if 'RESPONS' in limpa(val(cab - 1, c)).upper()), None)
    col_stat = next((c for c in range(1, ws.max_column + 1) if 'STATUS' in limpa(val(cab - 1, c)).upper()), None)
    plano = []
    for r in range(cab + 1, ws.max_row + 1):
        if limpa(val(r, 1)).upper().startswith('LEGENDA'):
            break
        oque = limpa(val(r, 3))
        if not oque:
            continue
        plano.append({
            'setor': limpa(val(r, 1)), 'ghe': limpa(val(r, 2)), 'oque': oque, 'porque': limpa(val(r, 4)),
            'como': limpa(val(r, 5)), 'onde': limpa(val(r, 6)), 'quanto': limpa(val(r, 7)), 'prioridade': limpa(val(r, 8)),
            'meses': [m for m, c in col_mes.items() if limpa(val(r, c))],
            'responsavel': limpa(val(r, col_resp)) if col_resp else '', 'status': limpa(val(r, col_stat)) if col_stat else '',
            'afericao': '',
        })
    return plano


if __name__ == '__main__':
    src, dst = sys.argv[1], sys.argv[2]
    wb = openpyxl.load_workbook(src, data_only=False)
    inv = ler_inventario(acha_aba(wb, 'INVENT'))
    pl = acha_aba(wb, 'PLANO')
    dados = {'empresa': {}, 'inventario': inv, 'plano': ler_plano(pl) if pl else []}
    json.dump(dados, open(dst, 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
    print(f'{len(inv)} linhas de risco, {len({(x["ghe"]) for x in inv})} GHE, {len(dados["plano"])} ações')
